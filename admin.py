"""
Admin module - Dashboard and post editor
"""
import os
import re
import base64
from io import BytesIO
from flask import Blueprint, render_template, request, redirect, url_for, flash, current_app, jsonify, send_file
from werkzeug.utils import secure_filename
from models import db, Post, User, Comment, Like, CommentLike, ReadEvent, MonthVisibility
from auth import admin_required
from datetime import datetime
from sqlalchemy import func
from sqlalchemy.orm import aliased
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from PIL import Image


def parse_user_agent(ua_string):
    """Parse user agent string into a readable browser/OS label."""
    if not ua_string:
        return 'Unknown'

    browser = 'Unknown Browser'
    os_name = 'Unknown OS'

    # Detect browser
    if 'Edg/' in ua_string:
        match = re.search(r'Edg/([\d.]+)', ua_string)
        browser = f"Edge {match.group(1).split('.')[0]}" if match else 'Edge'
    elif 'Chrome/' in ua_string and 'Safari/' in ua_string:
        match = re.search(r'Chrome/([\d.]+)', ua_string)
        browser = f"Chrome {match.group(1).split('.')[0]}" if match else 'Chrome'
    elif 'Firefox/' in ua_string:
        match = re.search(r'Firefox/([\d.]+)', ua_string)
        browser = f"Firefox {match.group(1).split('.')[0]}" if match else 'Firefox'
    elif 'Safari/' in ua_string and 'Chrome/' not in ua_string:
        match = re.search(r'Version/([\d.]+)', ua_string)
        browser = f"Safari {match.group(1).split('.')[0]}" if match else 'Safari'

    # Detect OS
    if 'Windows NT 10' in ua_string:
        os_name = 'Windows 10/11'
    elif 'Windows NT' in ua_string:
        os_name = 'Windows'
    elif 'Mac OS X' in ua_string:
        os_name = 'macOS'
    elif 'Linux' in ua_string and 'Android' not in ua_string:
        os_name = 'Linux'
    elif 'Android' in ua_string:
        os_name = 'Android'
    elif 'iPhone' in ua_string or 'iPad' in ua_string:
        os_name = 'iOS'

    return f"{browser} / {os_name}"

admin_bp = Blueprint('admin', __name__)


def allowed_file(filename):
    """Check if file extension is allowed"""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in current_app.config.get('ALLOWED_EXTENSIONS', {'png', 'jpg', 'jpeg', 'gif', 'webp'})


@admin_bp.route('/')
@admin_required
def dashboard():
    """Admin dashboard with overview stats"""
    # Get overall stats
    total_posts = Post.query.filter_by(status='published').count()
    total_users = User.query.count()
    total_comments = Comment.query.count()
    total_likes = Like.query.count()
    
    # Get recent posts with stats
    posts = Post.query.order_by(Post.created_at.desc()).limit(10).all()
    
    post_stats = []
    for post in posts:
        # Count unique viewers (logged-in + anonymous)
        unique_users = db.session.query(ReadEvent.user_id).filter(
            ReadEvent.post_id == post.id,
            ReadEvent.user_id.isnot(None)
        ).distinct().count()
        
        unique_anon = db.session.query(ReadEvent.anon_id).filter(
            ReadEvent.post_id == post.id,
            ReadEvent.anon_id.isnot(None)
        ).distinct().count()
        
        views = unique_users + unique_anon
        likes = Like.query.filter_by(post_id=post.id).count()
        comments = Comment.query.filter_by(post_id=post.id).count()
        
        # Calculate average completion
        avg_completion = db.session.query(func.avg(ReadEvent.percent)).filter_by(post_id=post.id).scalar() or 0
        
        post_stats.append({
            'post': post,
            'views': views,
            'likes': likes,
            'comments': comments,
            'avg_completion': round(avg_completion, 1)
        })
    
    return render_template('admin/dashboard.html', 
                          total_posts=total_posts,
                          total_users=total_users,
                          total_comments=total_comments,
                          total_likes=total_likes,
                          post_stats=post_stats)


@admin_bp.route('/analytics')
@admin_required
def analytics():
    """Redirect to latest available month analytics (or show empty state)."""
    latest_month = db.session.query(Post.month_key).filter_by(status='published').order_by(Post.month_key.desc()).first()
    if latest_month and latest_month[0]:
        return redirect(url_for('admin.month_analytics', month_key=latest_month[0]))

    months = db.session.query(Post.month_key).filter_by(status='published').distinct().order_by(Post.month_key.desc()).all()
    months = [m[0] for m in months if m and m[0]]
    flash('No published posts found yet to analyze.', 'warning')
    return render_template('admin/month_analytics.html', month_key=None, months=months, summary={}, posts=[], tables={})


@admin_bp.route('/analytics/month/<month_key>')
@admin_required
def month_analytics(month_key):
    """Monthly analytics dashboard based on Post.month_key."""
    # Month list for selector
    months = db.session.query(Post.month_key).filter_by(status='published').distinct().order_by(Post.month_key.desc()).all()
    months = [m[0] for m in months if m and m[0]]

    # Posts for month
    posts = Post.query.filter_by(status='published', month_key=month_key).order_by(Post.created_at.desc()).all()
    post_ids = [p.id for p in posts]

    if not posts:
        flash(f'No published posts found for {month_key}.', 'warning')
        return render_template(
            'admin/month_analytics.html',
            month_key=month_key,
            months=months,
            summary={
                'total_posts': 0,
                'total_reads': 0,
                'unique_readers': 0,
                'unique_logged_in_readers': 0,
                'unique_anon_readers': 0,
                'total_likes': 0,
                'total_comment_likes': 0,
                'total_comments': 0,
                'total_registered_users': User.query.count(),
            },
            posts=[],
            tables={},
        )

    # Summary metrics for month
    total_posts = len(posts)
    total_registered_users = User.query.count()

    total_reads = ReadEvent.query.filter(ReadEvent.post_id.in_(post_ids)).count()
    unique_logged_in_readers = db.session.query(ReadEvent.user_id).filter(
        ReadEvent.post_id.in_(post_ids),
        ReadEvent.user_id.isnot(None)
    ).distinct().count()
    unique_anon_readers = db.session.query(ReadEvent.anon_id).filter(
        ReadEvent.post_id.in_(post_ids),
        ReadEvent.anon_id.isnot(None)
    ).distinct().count()
    unique_readers = unique_logged_in_readers + unique_anon_readers

    total_likes = Like.query.filter(Like.post_id.in_(post_ids)).count()
    total_comments = Comment.query.filter(Comment.post_id.in_(post_ids)).count()
    total_comment_likes = (
        db.session.query(CommentLike.id)
        .join(Comment, Comment.id == CommentLike.comment_id)
        .join(Post, Post.id == Comment.post_id)
        .filter(Post.month_key == month_key, Post.status == 'published')
        .count()
    )

    # Aggregations per post
    likes_by_post = dict(
        db.session.query(Like.post_id, func.count(Like.id))
        .filter(Like.post_id.in_(post_ids))
        .group_by(Like.post_id)
        .all()
    )
    comments_by_post = dict(
        db.session.query(Comment.post_id, func.count(Comment.id))
        .filter(Comment.post_id.in_(post_ids))
        .group_by(Comment.post_id)
        .all()
    )
    comment_likes_by_comment = dict(
        db.session.query(CommentLike.comment_id, func.count(CommentLike.id))
        .join(Comment, Comment.id == CommentLike.comment_id)
        .filter(Comment.post_id.in_(post_ids))
        .group_by(CommentLike.comment_id)
        .all()
    )
    avg_completion_by_post = dict(
        db.session.query(ReadEvent.post_id, func.avg(ReadEvent.percent))
        .filter(ReadEvent.post_id.in_(post_ids))
        .group_by(ReadEvent.post_id)
        .all()
    )
    seconds_by_post = dict(
        db.session.query(ReadEvent.post_id, func.coalesce(func.sum(ReadEvent.seconds), 0))
        .filter(ReadEvent.post_id.in_(post_ids))
        .group_by(ReadEvent.post_id)
        .all()
    )

    post_rows = []
    for p in posts:
        # Unique viewers (logged-in + anon) per post
        post_unique_users = db.session.query(ReadEvent.user_id).filter(
            ReadEvent.post_id == p.id,
            ReadEvent.user_id.isnot(None)
        ).distinct().count()
        post_unique_anon = db.session.query(ReadEvent.anon_id).filter(
            ReadEvent.post_id == p.id,
            ReadEvent.anon_id.isnot(None)
        ).distinct().count()
        post_views = post_unique_users + post_unique_anon

        post_rows.append({
            'post': p,
            'views': post_views,
            'unique_users': post_unique_users,
            'unique_anon': post_unique_anon,
            'likes': int(likes_by_post.get(p.id, 0) or 0),
            'comments': int(comments_by_post.get(p.id, 0) or 0),
            'avg_completion': round(float(avg_completion_by_post.get(p.id, 0) or 0), 1),
            'total_seconds': int(seconds_by_post.get(p.id, 0) or 0),
        })

    # Sort helper tables
    top_by_views = sorted(post_rows, key=lambda r: r['views'], reverse=True)[:10]
    top_by_likes = sorted(post_rows, key=lambda r: r['likes'], reverse=True)[:10]
    top_by_comments = sorted(post_rows, key=lambda r: r['comments'], reverse=True)[:10]

    # Top comments by comment-like count (for posts in this month)
    month_comments = (
        Comment.query
        .filter(Comment.post_id.in_(post_ids))
        .order_by(Comment.created_at.desc())
        .all()
    )
    comment_rows = [{
        'comment': c,
        'post_id': c.post_id,
        'likes': int(comment_likes_by_comment.get(c.id, 0) or 0),
    } for c in month_comments]
    top_comments_by_likes = sorted(comment_rows, key=lambda r: r['likes'], reverse=True)[:20]

    # Who liked what (for posts in this month)
    like_events = (
        db.session.query(Like, User, Post)
        .join(User, User.id == Like.user_id)
        .join(Post, Post.id == Like.post_id)
        .filter(Post.month_key == month_key, Post.status == 'published')
        .order_by(Like.created_at.desc())
        .limit(300)
        .all()
    )
    like_rows = [{
        'liked_at': like.created_at,
        'user_name': user.name,
        'user_email': user.email,
        'post_title': post.title,
        'post_slug': post.slug,
    } for (like, user, post) in like_events]

    # Who commented what (for posts in this month)
    comment_events = (
        db.session.query(Comment, User, Post)
        .join(User, User.id == Comment.user_id)
        .join(Post, Post.id == Comment.post_id)
        .filter(Post.month_key == month_key, Post.status == 'published')
        .order_by(Comment.created_at.desc())
        .limit(300)
        .all()
    )
    comment_event_rows = [{
        'commented_at': c.created_at,
        'user_name': u.name,
        'user_email': u.email,
        'post_title': p.title,
        'post_slug': p.slug,
        'body': (c.body or '')[:220],
    } for (c, u, p) in comment_events]

    # Who liked which comment (for posts in this month)
    Liker = aliased(User)
    Author = aliased(User)
    comment_like_events = (
        db.session.query(CommentLike, Liker, Comment, Post, Author)
        .join(Liker, Liker.id == CommentLike.user_id)
        .join(Comment, Comment.id == CommentLike.comment_id)
        .join(Post, Post.id == Comment.post_id)
        .join(Author, Author.id == Comment.user_id)
        .filter(Post.month_key == month_key, Post.status == 'published')
        .order_by(CommentLike.created_at.desc())
        .limit(300)
        .all()
    )
    comment_like_rows = [{
        'liked_at': cl.created_at,
        'liker_name': liker.name,
        'liker_email': liker.email,
        'author_name': author.name,
        'author_email': author.email,
        'post_title': post.title,
        'post_slug': post.slug,
        'comment_body': (comment.body or '')[:220],
    } for (cl, liker, comment, post, author) in comment_like_events]

    # "Where" (reads only, since Like/Comment don't store ip/user_agent yet)
    top_ips = (
        db.session.query(ReadEvent.ip_address, func.count(ReadEvent.id))
        .filter(ReadEvent.post_id.in_(post_ids), ReadEvent.ip_address.isnot(None))
        .group_by(ReadEvent.ip_address)
        .order_by(func.count(ReadEvent.id).desc())
        .limit(20)
        .all()
    )
    top_user_agents_raw = (
        db.session.query(ReadEvent.user_agent, func.count(ReadEvent.id))
        .filter(ReadEvent.post_id.in_(post_ids), ReadEvent.user_agent.isnot(None))
        .group_by(ReadEvent.user_agent)
        .order_by(func.count(ReadEvent.id).desc())
        .limit(20)
        .all()
    )
    # Parse user agents into readable labels and aggregate
    ua_parsed = {}
    for ua, cnt in top_user_agents_raw:
        label = parse_user_agent(ua)
        ua_parsed[label] = ua_parsed.get(label, 0) + int(cnt)
    top_user_agents = sorted(ua_parsed.items(), key=lambda x: x[1], reverse=True)[:15]

    recent_reads = (
        ReadEvent.query
        .filter(ReadEvent.post_id.in_(post_ids))
        .order_by(ReadEvent.created_at.desc())
        .limit(200)
        .all()
    )

    # Previous month comparison
    prev_month_key = None
    prev_summary = {}
    try:
        year, month = int(month_key[:4]), int(month_key[5:7])
        if month == 1:
            prev_month_key = f"{year - 1}-12"
        else:
            prev_month_key = f"{year}-{month - 1:02d}"

        prev_post_ids = [p.id for p in Post.query.filter_by(status='published', month_key=prev_month_key).all()]
        if prev_post_ids:
            prev_summary = {
                'total_reads': ReadEvent.query.filter(ReadEvent.post_id.in_(prev_post_ids)).count(),
                'unique_readers': (
                    db.session.query(ReadEvent.user_id).filter(ReadEvent.post_id.in_(prev_post_ids), ReadEvent.user_id.isnot(None)).distinct().count() +
                    db.session.query(ReadEvent.anon_id).filter(ReadEvent.post_id.in_(prev_post_ids), ReadEvent.anon_id.isnot(None)).distinct().count()
                ),
                'total_likes': Like.query.filter(Like.post_id.in_(prev_post_ids)).count(),
                'total_comments': Comment.query.filter(Comment.post_id.in_(prev_post_ids)).count(),
            }
    except Exception:
        pass

    # Reads by day (micro granularity)
    reads_by_day = (
        db.session.query(func.date(ReadEvent.created_at), func.count(ReadEvent.id))
        .filter(ReadEvent.post_id.in_(post_ids))
        .group_by(func.date(ReadEvent.created_at))
        .order_by(func.date(ReadEvent.created_at).asc())
        .all()
    )
    reads_by_day_rows = [{'day': d, 'reads': int(c)} for (d, c) in reads_by_day]

    # Reads by hour (what time people are reading)
    reads_by_hour_raw = (
        db.session.query(func.extract('hour', ReadEvent.created_at).label('hr'), func.count(ReadEvent.id))
        .filter(ReadEvent.post_id.in_(post_ids))
        .group_by('hr')
        .order_by('hr')
        .all()
    )
    hour_map = {int(hr): int(cnt) for (hr, cnt) in reads_by_hour_raw if hr is not None}
    reads_by_hour_rows = [{'hour': h, 'reads': int(hour_map.get(h, 0))} for h in range(24)]

    # Reads by day-of-week (0=Sunday..6=Saturday in Postgres)
    reads_by_dow_raw = (
        db.session.query(func.extract('dow', ReadEvent.created_at).label('dow'), func.count(ReadEvent.id))
        .filter(ReadEvent.post_id.in_(post_ids))
        .group_by('dow')
        .order_by('dow')
        .all()
    )
    dow_names = {0: 'Sun', 1: 'Mon', 2: 'Tue', 3: 'Wed', 4: 'Thu', 5: 'Fri', 6: 'Sat'}
    dow_map = {int(d): int(cnt) for (d, cnt) in reads_by_dow_raw if d is not None}
    reads_by_dow_rows = [{'dow': d, 'day': dow_names.get(d, str(d)), 'reads': int(dow_map.get(d, 0))} for d in range(7)]

    # Scroll/reading progress distribution (how they are moving)
    percent_values = (
        db.session.query(ReadEvent.percent)
        .filter(ReadEvent.post_id.in_(post_ids), ReadEvent.percent.isnot(None))
        .all()
    )
    buckets = [
        ('0-10%', 0, 10),
        ('11-25%', 11, 25),
        ('26-50%', 26, 50),
        ('51-75%', 51, 75),
        ('76-90%', 76, 90),
        ('91-100%', 91, 100),
    ]
    bucket_counts = {label: 0 for (label, _, _) in buckets}
    for (p,) in percent_values:
        try:
            pv = int(p)
        except Exception:
            continue
        for (label, lo, hi) in buckets:
            if lo <= pv <= hi:
                bucket_counts[label] += 1
                break
    progress_buckets = [{'label': label, 'count': int(bucket_counts[label])} for (label, _, _) in buckets]

    # Post-to-post transitions (approximate "moving" between posts)
    # We infer a transition if the same reader (user_id or anon_id) reads different posts within 30 minutes.
    events = (
        db.session.query(ReadEvent.created_at, ReadEvent.post_id, ReadEvent.user_id, ReadEvent.anon_id)
        .filter(ReadEvent.post_id.in_(post_ids))
        .order_by(ReadEvent.user_id.asc().nulls_last(), ReadEvent.anon_id.asc().nulls_last(), ReadEvent.created_at.asc())
        .limit(20000)
        .all()
    )
    def reader_key(user_id, anon_id):
        if user_id is not None:
            return f"u:{user_id}"
        if anon_id:
            return f"a:{anon_id}"
        return None

    from datetime import timedelta
    transition_counts = {}
    last_by_reader = {}
    for created_at, post_id, user_id, anon_id in events:
        key = reader_key(user_id, anon_id)
        if not key or not created_at or not post_id:
            continue
        last = last_by_reader.get(key)
        if last:
            last_time, last_post = last
            if last_post != post_id and (created_at - last_time) <= timedelta(minutes=30):
                transition_counts[(last_post, post_id)] = transition_counts.get((last_post, post_id), 0) + 1
        last_by_reader[key] = (created_at, post_id)

    post_title_by_id = {p.id: p.title for p in posts}
    post_slug_by_id = {p.id: p.slug for p in posts}
    transitions_top = sorted(transition_counts.items(), key=lambda kv: kv[1], reverse=True)[:25]
    transition_rows = [{
        'from_post_id': a,
        'to_post_id': b,
        'from_title': post_title_by_id.get(a, f'Post {a}'),
        'to_title': post_title_by_id.get(b, f'Post {b}'),
        'from_slug': post_slug_by_id.get(a),
        'to_slug': post_slug_by_id.get(b),
        'count': int(cnt),
    } for ((a, b), cnt) in transitions_top]

    summary = {
        'total_posts': total_posts,
        'total_reads': total_reads,
        'unique_readers': unique_readers,
        'unique_logged_in_readers': unique_logged_in_readers,
        'unique_anon_readers': unique_anon_readers,
        'total_likes': total_likes,
        'total_comment_likes': total_comment_likes,
        'total_comments': total_comments,
        'total_registered_users': total_registered_users,
    }

    tables = {
        'post_rows': post_rows,
        'top_by_views': top_by_views,
        'top_by_likes': top_by_likes,
        'top_by_comments': top_by_comments,
        'like_rows': like_rows,
        'comment_rows': comment_event_rows,
        'comment_like_rows': comment_like_rows,
        'top_comments_by_likes': top_comments_by_likes,
        'top_ips': top_ips,
        'top_user_agents': top_user_agents,
        'recent_reads': recent_reads,
        'reads_by_day': reads_by_day_rows,
        'reads_by_hour': reads_by_hour_rows,
        'reads_by_dow': reads_by_dow_rows,
        'progress_buckets': progress_buckets,
        'transitions': transition_rows,
    }

    return render_template(
        'admin/month_analytics.html',
        month_key=month_key,
        months=months,
        summary=summary,
        prev_summary=prev_summary,
        posts=posts,
        tables=tables,
    )


@admin_bp.route('/posts')
@admin_required
def posts_list():
    """List all posts"""
    posts = Post.query.order_by(Post.created_at.desc()).all()
    return render_template('admin/posts_list.html', posts=posts)


@admin_bp.route('/posts/new', methods=['GET', 'POST'])
@admin_required
def post_new():
    """Create new post"""
    if request.method == 'POST':
        try:
            title = request.form.get('title', '').strip()
            slug = request.form.get('slug', '').strip()
            month_key = request.form.get('month_key', '').strip()
            status = request.form.get('status', 'draft')
            html_content = request.form.get('html_content', '')
            excerpt = request.form.get('excerpt', '').strip()
            category = request.form.get('category', '').strip()
            read_time = request.form.get('read_time', 0)
            is_featured = request.form.get('is_featured') == 'on'  # Checkbox value
            
            # Validate required fields
            if not title or not slug or not month_key:
                flash('Title, slug, and month are required', 'error')
                return redirect(url_for('admin.post_new'))
            
            # Check if slug exists
            existing = Post.query.filter_by(slug=slug).first()
            if existing:
                flash('Slug already exists', 'error')
                return redirect(url_for('admin.post_new'))
            
            # Handle published date
            published_at = None
            if status == 'published':
                published_date = request.form.get('published_date')
                published_time = request.form.get('published_time', '00:00')
                if published_date:
                    published_at = datetime.strptime(f"{published_date} {published_time}", "%Y-%m-%d %H:%M")
                else:
                    published_at = datetime.utcnow()
            
            # Handle hero image upload
            hero_image_path = None
            if 'hero_image' in request.files:
                file = request.files['hero_image']
                if file and file.filename and allowed_file(file.filename):
                    filename = secure_filename(file.filename)
                    # Add timestamp to avoid conflicts
                    filename = f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{filename}"
                    filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
                    file.save(filepath)
                    # Use forward slashes for web URLs
                    hero_image_path = f"/{filepath.replace(chr(92), '/')}"  # chr(92) is backslash
            
            # Handle related posts
            related_posts = request.form.getlist('related_posts')
            
            # Create post
            post = Post(
                title=title,
                slug=slug,
                month_key=month_key,
                status=status,
                html_content=html_content,
                excerpt=excerpt,
                category=category,
                read_time=read_time,
                is_featured=is_featured,
                hero_image_path=hero_image_path,
                published_at=published_at
            )
            
            # Set related posts
            if related_posts:
                post.set_related_posts(related_posts)
            
            db.session.add(post)
            db.session.commit()

            # Auto-create MonthVisibility record if it doesn't exist
            existing_mv = MonthVisibility.query.filter_by(month_key=month_key).first()
            if not existing_mv:
                mv = MonthVisibility(month_key=month_key, visibility='prerelease')
                db.session.add(mv)
                db.session.commit()

            flash('Post created successfully', 'success')
            return redirect(url_for('admin.posts_list'))

        except Exception as e:
            current_app.logger.error(f"Post creation error: {e}")
            if '413' in str(e) or 'Too Large' in str(e):
                flash('Failed to create post: Content exceeds the 16MB size limit. Please reduce image sizes or use fewer images.', 'error')
            else:
                flash('Failed to create post', 'error')
            return redirect(url_for('admin.post_new'))
    
    all_posts = Post.query.filter_by(status='published').order_by(Post.published_at.desc()).all()
    return render_template('admin/post_edit.html', post=None, all_posts=all_posts)


@admin_bp.route('/posts/<int:post_id>/edit', methods=['GET', 'POST'])
@admin_required
def post_edit(post_id):
    """Edit existing post"""
    post = Post.query.get_or_404(post_id)
    
    if request.method == 'POST':
        try:
            post.title = request.form.get('title', '').strip()
            post.slug = request.form.get('slug', '').strip()
            post.month_key = request.form.get('month_key', '').strip()
            post.status = request.form.get('status', 'draft')
            post.html_content = request.form.get('html_content', '')
            post.excerpt = request.form.get('excerpt', '').strip()
            post.category = request.form.get('category', '').strip()
            post.read_time = request.form.get('read_time', 0)
            post.is_featured = request.form.get('is_featured') == 'on'  # Checkbox value
            
            # Debug logging
            print(f"📝 [DEBUG] Received html_content with {len(post.html_content)} characters")
            print(f"📝 [DEBUG] Content preview: {post.html_content[:200] if post.html_content else 'EMPTY'}")
            print(f"📝 [DEBUG] Post title: {post.title}")
            print(f"📝 [DEBUG] Post slug: {post.slug}")
            
            # Validate required fields
            if not post.title or not post.slug or not post.month_key:
                flash('Title, slug, and month are required', 'error')
                return redirect(url_for('admin.post_edit', post_id=post_id))
            
            # Check if slug exists (excluding current post)
            existing = Post.query.filter(Post.slug == post.slug, Post.id != post_id).first()
            if existing:
                flash('Slug already exists', 'error')
                return redirect(url_for('admin.post_edit', post_id=post_id))
            
            # Handle published date
            if post.status == 'published':
                published_date = request.form.get('published_date')
                published_time = request.form.get('published_time', '00:00')
                if published_date:
                    post.published_at = datetime.strptime(f"{published_date} {published_time}", "%Y-%m-%d %H:%M")
                elif not post.published_at:
                    post.published_at = datetime.utcnow()
            
            # Handle hero image upload
            if 'hero_image' in request.files:
                file = request.files['hero_image']
                if file and file.filename and allowed_file(file.filename):
                    filename = secure_filename(file.filename)
                    filename = f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{filename}"
                    filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
                    file.save(filepath)
                    # Use forward slashes for web URLs
                    post.hero_image_path = f"/{filepath.replace(chr(92), '/')}"  # chr(92) is backslash
            
            # Handle related posts
            related_posts = request.form.getlist('related_posts')
            post.set_related_posts(related_posts if related_posts else None)
            
            post.updated_at = datetime.utcnow()

            db.session.commit()

            # Auto-create MonthVisibility record if month_key changed
            existing_mv = MonthVisibility.query.filter_by(month_key=post.month_key).first()
            if not existing_mv:
                mv = MonthVisibility(month_key=post.month_key, visibility='prerelease')
                db.session.add(mv)
                db.session.commit()

            flash('Post updated successfully', 'success')
            return redirect(url_for('admin.posts_list'))
            
        except Exception as e:
            current_app.logger.error(f"Post update error: {e}")
            if '413' in str(e) or 'Too Large' in str(e):
                flash('Failed to update post: Content exceeds the 16MB size limit. Please reduce image sizes or use fewer images.', 'error')
            else:
                flash('Failed to update post', 'error')
            return redirect(url_for('admin.post_edit', post_id=post_id))
    
    all_posts = Post.query.filter_by(status='published').order_by(Post.published_at.desc()).all()
    return render_template('admin/post_edit.html', post=post, all_posts=all_posts)


@admin_bp.route('/posts/<int:post_id>/delete', methods=['POST'])
@admin_required
def post_delete(post_id):
    """Delete a post"""
    try:
        post = Post.query.get_or_404(post_id)
        db.session.delete(post)
        db.session.commit()
        flash('Post deleted successfully', 'success')
    except Exception as e:
        current_app.logger.error(f"Post delete error: {e}")
        flash('Failed to delete post', 'error')
    
    return redirect(url_for('admin.posts_list'))


@admin_bp.route('/posts/<int:post_id>/stats')
@admin_required
def post_stats(post_id):
    """Detailed stats for a specific post"""
    from models import CommentLike
    from sqlalchemy import func
    
    post = Post.query.get_or_404(post_id)
    
    # Get unique viewers (logged-in + anonymous)
    # Count distinct user_id (for logged-in) and anon_id (for anonymous)
    unique_users = db.session.query(ReadEvent.user_id).filter(
        ReadEvent.post_id == post_id,
        ReadEvent.user_id.isnot(None)
    ).distinct().count()
    
    unique_anon = db.session.query(ReadEvent.anon_id).filter(
        ReadEvent.post_id == post_id,
        ReadEvent.anon_id.isnot(None)
    ).distinct().count()
    
    total_views = unique_users + unique_anon
    
    # Get logged-in viewers
    viewers = db.session.query(User).join(ReadEvent).filter(ReadEvent.post_id == post_id).distinct().all()

    # Get anonymous reader summaries
    anon_readers = db.session.query(
        ReadEvent.anon_id,
        func.count(ReadEvent.id).label('events'),
        func.max(ReadEvent.created_at).label('last_seen'),
        func.max(ReadEvent.ip_address).label('ip_address')
    ).filter(
        ReadEvent.post_id == post_id,
        ReadEvent.anon_id.isnot(None)
    ).group_by(ReadEvent.anon_id).all()

    # Get all read events
    read_events = ReadEvent.query.filter_by(post_id=post_id).order_by(ReadEvent.created_at.desc()).all()
    avg_completion = db.session.query(func.avg(ReadEvent.percent)).filter_by(post_id=post_id).scalar() or 0
    avg_time = db.session.query(func.avg(ReadEvent.seconds)).filter_by(post_id=post_id).scalar() or 0
    
    # Get likes
    likes = Like.query.filter_by(post_id=post_id).all()
    
    # Get all comments (including nested)
    all_comments = Comment.query.filter_by(post_id=post_id).all()
    top_level_comments = [c for c in all_comments if c.parent_id is None]
    reply_comments = [c for c in all_comments if c.parent_id is not None]
    
    # Get comment likes for this post
    comment_ids = [c.id for c in all_comments]
    comment_likes_count = CommentLike.query.filter(CommentLike.comment_id.in_(comment_ids)).count() if comment_ids else 0
    
    # Build comment data with nested structure
    comment_data = []
    for comment in sorted(top_level_comments, key=lambda x: x.created_at, reverse=True):
        data = {
            'comment': comment,
            'user': comment.user,
            'like_count': comment.get_like_count(),
            'reply_count': comment.get_reply_count(),
            'replies': []
        }
        
        # Get replies
        for reply in comment.get_all_replies():
            reply_data = {
                'comment': reply,
                'user': reply.user,
                'like_count': reply.get_like_count()
            }
            data['replies'].append(reply_data)
        
        comment_data.append(data)
    
    # Calculate engagement rate
    engagement_rate = 0
    if total_views > 0:
        total_interactions = len(likes) + len(all_comments) + comment_likes_count
        engagement_rate = round((total_interactions / total_views) * 100, 1)
    
    return render_template('admin/post_stats.html',
                          post=post,
                          viewers=viewers,
                          anon_readers=anon_readers,
                          total_views=total_views,
                          avg_completion=round(avg_completion, 1),
                          avg_time=round(avg_time / 60, 1) if avg_time else 0,
                          likes=likes,
                          read_events=read_events,
                          comment_data=comment_data,
                          total_comments=len(all_comments),
                          top_level_count=len(top_level_comments),
                          reply_count=len(reply_comments),
                          comment_likes_count=comment_likes_count,
                          engagement_rate=engagement_rate)


@admin_bp.route('/posts/<int:post_id>/readers')
@admin_required
def post_readers(post_id):
    """Detailed reader log for a specific post"""
    post = Post.query.get_or_404(post_id)
    
    # Get pagination parameters
    page = request.args.get('page', 1, type=int)
    per_page = 50
    
    # Get filters
    reader_type = request.args.get('reader_type', 'all')  # all, logged_in, anonymous
    
    # Build query
    query = ReadEvent.query.filter_by(post_id=post_id).order_by(ReadEvent.created_at.desc())
    
    if reader_type == 'logged_in':
        query = query.filter(ReadEvent.user_id.isnot(None))
    elif reader_type == 'anonymous':
        query = query.filter(ReadEvent.anon_id.isnot(None))
    
    # Paginate
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    read_events = pagination.items
    
    # Calculate stats
    total_events = query.count()
    unique_users = db.session.query(ReadEvent.user_id).filter(
        ReadEvent.post_id == post_id,
        ReadEvent.user_id.isnot(None)
    ).distinct().count()
    
    unique_anon = db.session.query(ReadEvent.anon_id).filter(
        ReadEvent.post_id == post_id,
        ReadEvent.anon_id.isnot(None)
    ).distinct().count()
    
    return render_template('admin/post_readers.html',
                          post=post,
                          read_events=read_events,
                          pagination=pagination,
                          total_events=total_events,
                          unique_users=unique_users,
                          unique_anon=unique_anon,
                          reader_type=reader_type)


@admin_bp.route('/users')
@admin_required
def users_list():
    """List all users"""
    users = User.query.order_by(User.last_login_at.desc()).all()
    return render_template('admin/users_list.html', users=users)


@admin_bp.route('/readers')
@admin_required
def readers_list():
    """Global readers view with aggregation"""
    from sqlalchemy import func, case
    
    # Get pagination
    page = request.args.get('page', 1, type=int)
    per_page = 50
    
    # Get filters
    post_filter = request.args.get('post_id', type=int)
    
    # Build reader aggregation query
    # We need to group by either user_id or anon_id
    
    # Get logged-in users with stats
    logged_in_readers = db.session.query(
        User.id.label('user_id'),
        User.email.label('email'),
        User.name.label('name'),
        func.count(func.distinct(ReadEvent.post_id)).label('posts_read'),
        func.count(ReadEvent.id).label('total_events'),
        func.max(ReadEvent.created_at).label('last_seen'),
        func.min(ReadEvent.created_at).label('first_seen')
    ).join(ReadEvent, ReadEvent.user_id == User.id)
    
    if post_filter:
        logged_in_readers = logged_in_readers.filter(ReadEvent.post_id == post_filter)
    
    logged_in_readers = logged_in_readers.group_by(User.id, User.email, User.name).all()
    
    # Get anonymous readers with stats
    # Subquery to get the IP from the most recent event per anon_id
    latest_anon_event = db.session.query(
        ReadEvent.anon_id,
        func.max(ReadEvent.id).label('latest_id')
    ).filter(ReadEvent.anon_id.isnot(None)).group_by(ReadEvent.anon_id).subquery()

    latest_ip = db.session.query(
        latest_anon_event.c.anon_id,
        ReadEvent.ip_address
    ).join(ReadEvent, ReadEvent.id == latest_anon_event.c.latest_id).subquery()

    anon_readers = db.session.query(
        ReadEvent.anon_id.label('anon_id'),
        func.count(func.distinct(ReadEvent.post_id)).label('posts_read'),
        func.count(ReadEvent.id).label('total_events'),
        func.max(ReadEvent.created_at).label('last_seen'),
        func.min(ReadEvent.created_at).label('first_seen'),
        latest_ip.c.ip_address.label('ip_address')
    ).outerjoin(latest_ip, latest_ip.c.anon_id == ReadEvent.anon_id).filter(ReadEvent.anon_id.isnot(None))
    
    if post_filter:
        anon_readers = anon_readers.filter(ReadEvent.post_id == post_filter)
    
    anon_readers = anon_readers.group_by(ReadEvent.anon_id, latest_ip.c.ip_address).all()
    
    # Combine and format
    readers = []
    
    for reader in logged_in_readers:
        readers.append({
            'type': 'logged_in',
            'user_id': reader.user_id,
            'label': reader.name or reader.email,
            'email': reader.email,
            'posts_read': reader.posts_read,
            'total_events': reader.total_events,
            'last_seen': reader.last_seen,
            'first_seen': reader.first_seen,
            'ip_address': None
        })
    
    for reader in anon_readers:
        readers.append({
            'type': 'anonymous',
            'anon_id': reader.anon_id,
            'label': f'Anon ({reader.anon_id[:8]}...)',
            'posts_read': reader.posts_read,
            'total_events': reader.total_events,
            'last_seen': reader.last_seen,
            'first_seen': reader.first_seen,
            'ip_address': reader.ip_address
        })
    
    # Sort by last_seen desc
    readers.sort(key=lambda x: x['last_seen'], reverse=True)
    
    # Manual pagination
    total = len(readers)
    start = (page - 1) * per_page
    end = start + per_page
    readers_page = readers[start:end]
    
    # Get all posts for filter dropdown
    all_posts = Post.query.filter_by(status='published').order_by(Post.published_at.desc()).all()
    
    return render_template('admin/readers.html',
                          readers=readers_page,
                          page=page,
                          per_page=per_page,
                          total=total,
                          all_posts=all_posts,
                          post_filter=post_filter)


@admin_bp.route('/readers/<reader_type>/<reader_id>')
@admin_required
def reader_detail(reader_type, reader_id):
    """Detailed view of what a specific reader has read"""
    if reader_type == 'user':
        user = User.query.get_or_404(reader_id)
        reader_label = user.name or user.email
        read_events = ReadEvent.query.filter_by(user_id=reader_id).order_by(ReadEvent.created_at.desc()).all()
    else:  # anonymous
        reader_label = f'Anonymous ({reader_id[:8]}...)'
        read_events = ReadEvent.query.filter_by(anon_id=reader_id).order_by(ReadEvent.created_at.desc()).all()
    
    # Group by post
    posts_read = {}
    for event in read_events:
        if event.post_id not in posts_read:
            posts_read[event.post_id] = {
                'post': event.post,
                'events': [],
                'max_percent': 0,
                'total_time': 0
            }
        posts_read[event.post_id]['events'].append(event)
        posts_read[event.post_id]['max_percent'] = max(posts_read[event.post_id]['max_percent'], event.percent or 0)
        posts_read[event.post_id]['total_time'] += event.seconds or 0
    
    return render_template('admin/reader_detail.html',
                          reader_type=reader_type,
                          reader_id=reader_id,
                          reader_label=reader_label,
                          posts_read=posts_read,
                          total_events=len(read_events))


@admin_bp.route('/comments')
@admin_required
def comments_list():
    """List all comments with nested structure and moderation options"""
    from models import Comment, CommentLike
    
    # Get all top-level comments (no parent)
    comments = Comment.query.filter_by(parent_id=None).order_by(Comment.created_at.desc()).all()
    
    # Build comment data with stats
    comment_data = []
    for comment in comments:
        data = {
            'comment': comment,
            'user': comment.user,
            'post': comment.post,
            'like_count': comment.get_like_count(),
            'reply_count': comment.get_reply_count(),
            'replies': []
        }
        
        # Get all replies
        for reply in comment.get_all_replies():
            reply_data = {
                'comment': reply,
                'user': reply.user,
                'like_count': reply.get_like_count()
            }
            data['replies'].append(reply_data)
        
        comment_data.append(data)
    
    # Get stats
    total_comments = Comment.query.count()
    total_likes = CommentLike.query.count()
    
    return render_template('admin/comments_list.html', 
                         comment_data=comment_data,
                         total_comments=total_comments,
                         total_likes=total_likes)


@admin_bp.route('/comments/<int:comment_id>/delete', methods=['POST'])
@admin_required
def delete_comment_admin(comment_id):
    """Delete a comment (admin action)"""
    try:
        from models import Comment
        comment = Comment.query.get_or_404(comment_id)
        db.session.delete(comment)
        db.session.commit()
        flash('Comment deleted successfully', 'success')
    except Exception as e:
        flash(f'Error deleting comment: {str(e)}', 'error')
    
    return redirect(url_for('admin.comments_list'))


@admin_bp.route('/upload-image', methods=['POST'])
@admin_required
def upload_image():
    """Upload image for editor"""
    try:
        if 'image' not in request.files:
            return jsonify({'error': 'No image provided'}), 400
        
        file = request.files['image']
        if file.filename == '':
            return jsonify({'error': 'No file selected'}), 400
        
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            filename = f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{filename}"
            filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            
            # Use forward slashes for web URLs
            web_path = f"/{filepath.replace(chr(92), '/')}"  # chr(92) is backslash
            
            return jsonify({
                'success': True,
                'url': web_path
            })
        
        return jsonify({'error': 'Invalid file type'}), 400
        
    except Exception as e:
        current_app.logger.error(f"Image upload error: {e}")
        if '413' in str(e) or 'Too Large' in str(e):
            return jsonify({'error': 'Image exceeds the 16MB size limit. Please reduce the image size.'}), 413
        return jsonify({'error': 'Upload failed'}), 500


@admin_bp.route('/posts/export')
@admin_required
def posts_export():
    """Export all posts to Excel with images as base64"""
    try:
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Posts"
        
        # Define headers
        headers = [
            'ID', 'Slug', 'Title', 'Month Key', 'Published At', 
            'Status', 'Is Featured', 'Hero Image Path', 'Hero Image (Base64)', 'Hero Image Filename',
            'HTML Content', 'Excerpt', 'Category', 'Read Time', 'Created At', 'Updated At'
        ]
        
        # Write headers with styling
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Get all posts
        posts = Post.query.order_by(Post.created_at.desc()).all()
        
        # Write post data
        for row_num, post in enumerate(posts, 2):
            # Convert hero image to base64 if exists
            hero_image_base64 = ''
            hero_image_filename = ''
            
            if post.hero_image_path:
                # Try multiple possible image locations
                image_found = False
                possible_paths = [
                    # Try uploads folder first
                    os.path.join(current_app.config['UPLOAD_FOLDER'], os.path.basename(post.hero_image_path)),
                    # Try relative to app root
                    os.path.join('/app', post.hero_image_path.lstrip('/')),
                    # Try as absolute path
                    post.hero_image_path
                ]
                
                for image_path in possible_paths:
                    if os.path.exists(image_path):
                        try:
                            with open(image_path, 'rb') as img_file:
                                hero_image_base64 = base64.b64encode(img_file.read()).decode('utf-8')
                                hero_image_filename = os.path.basename(post.hero_image_path)
                                image_found = True
                                break
                        except Exception as e:
                            current_app.logger.error(f"Error reading image {image_path}: {e}")
                            continue
                
                if not image_found:
                    current_app.logger.warning(f"Image not found for post {post.slug}: {post.hero_image_path}")
            
            # Write row data
            ws.cell(row=row_num, column=1, value=post.id)
            ws.cell(row=row_num, column=2, value=post.slug)
            ws.cell(row=row_num, column=3, value=post.title)
            ws.cell(row=row_num, column=4, value=post.month_key)
            ws.cell(row=row_num, column=5, value=str(post.published_at) if post.published_at else '')
            ws.cell(row=row_num, column=6, value=post.status)
            ws.cell(row=row_num, column=7, value='Yes' if post.is_featured else 'No')
            ws.cell(row=row_num, column=8, value=post.hero_image_path or '')
            ws.cell(row=row_num, column=9, value=hero_image_base64)
            ws.cell(row=row_num, column=10, value=hero_image_filename)
            ws.cell(row=row_num, column=11, value=post.html_content or '')
            ws.cell(row=row_num, column=12, value=post.excerpt or '')
            ws.cell(row=row_num, column=13, value=post.category or '')
            ws.cell(row=row_num, column=14, value=post.read_time or '')
            ws.cell(row=row_num, column=15, value=str(post.created_at))
            ws.cell(row=row_num, column=16, value=str(post.updated_at) if post.updated_at else '')
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 10   # ID
        ws.column_dimensions['B'].width = 30   # Slug
        ws.column_dimensions['C'].width = 50   # Title
        ws.column_dimensions['D'].width = 15   # Month Key
        ws.column_dimensions['E'].width = 20   # Published At
        ws.column_dimensions['F'].width = 12   # Status
        ws.column_dimensions['G'].width = 12   # Is Featured
        ws.column_dimensions['H'].width = 40   # Hero Image Path
        ws.column_dimensions['I'].width = 20   # Base64 (collapsed)
        ws.column_dimensions['J'].width = 30   # Filename
        ws.column_dimensions['K'].width = 20   # HTML Content (collapsed)
        ws.column_dimensions['L'].width = 40   # Excerpt
        ws.column_dimensions['M'].width = 20   # Category
        ws.column_dimensions['N'].width = 12   # Read Time
        ws.column_dimensions['O'].width = 20   # Created At
        ws.column_dimensions['P'].width = 20   # Updated At
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Generate filename with timestamp
        filename = f"posts_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        # Don't flash message - it will show on the download page
        # flash(f'Successfully exported {len(posts)} posts to Excel', 'success')
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        current_app.logger.error(f"Export error: {e}")
        flash(f'Export failed: {str(e)}', 'error')
        return redirect(url_for('admin.posts_list'))


@admin_bp.route('/posts/import', methods=['GET', 'POST'])
@admin_required
def posts_import():
    """Import posts from Excel file"""
    if request.method == 'POST':
        try:
            # Check if file was uploaded
            if 'file' not in request.files:
                flash('No file uploaded', 'error')
                return redirect(url_for('admin.posts_list'))
            
            file = request.files['file']
            
            if file.filename == '':
                flash('No file selected', 'error')
                return redirect(url_for('admin.posts_list'))
            
            if not file.filename.endswith('.xlsx'):
                flash('Only .xlsx files are supported', 'error')
                return redirect(url_for('admin.posts_list'))
            
            # Load workbook
            wb = load_workbook(file, data_only=True)
            ws = wb.active
            
            # Get headers (first row)
            headers = [cell.value for cell in ws[1]]
            
            # Validate headers
            required_headers = ['Slug', 'Title', 'Month Key', 'Status']
            for header in required_headers:
                if header not in headers:
                    flash(f'Missing required column: {header}', 'error')
                    return redirect(url_for('admin.posts_list'))
            
            # Import posts
            imported_count = 0
            updated_count = 0
            skipped_count = 0
            
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                try:
                    # Create dict from row
                    row_data = dict(zip(headers, row))
                    
                    # Skip empty rows
                    if not row_data.get('Slug') or not row_data.get('Title'):
                        skipped_count += 1
                        continue
                    
                    # Check if post exists
                    existing_post = Post.query.filter_by(slug=row_data['Slug']).first()
                    
                    if existing_post:
                        # Update existing post
                        post = existing_post
                        updated_count += 1
                    else:
                        # Create new post
                        post = Post()
                        imported_count += 1
                    
                    # Set basic fields
                    post.slug = row_data['Slug']
                    post.title = row_data['Title']
                    post.month_key = row_data['Month Key']
                    post.status = row_data.get('Status', 'draft')
                    post.is_featured = row_data.get('Is Featured', 'No') == 'Yes'
                    post.html_content = row_data.get('HTML Content', '')
                    post.excerpt = row_data.get('Excerpt', '')
                    post.category = row_data.get('Category', '')
                    post.read_time = row_data.get('Read Time')
                    
                    # Parse published_at
                    if row_data.get('Published At'):
                        try:
                            post.published_at = datetime.strptime(str(row_data['Published At']), '%Y-%m-%d %H:%M:%S')
                        except:
                            # Try without time
                            try:
                                post.published_at = datetime.strptime(str(row_data['Published At']).split()[0], '%Y-%m-%d')
                            except:
                                pass
                    
                    # Handle hero image
                    # Priority: 1. Hero Image Path (if exists), 2. Base64 decode and save
                    hero_image_path = str(row_data.get('Hero Image Path', '') or '').strip()
                    hero_image_base64 = row_data.get('Hero Image (Base64)', '') or ''
                    hero_image_filename = row_data.get('Hero Image Filename', '') or ''
                    
                    if hero_image_path:
                        # Use existing path (for cases where image already exists in assets)
                        post.hero_image_path = hero_image_path
                    elif hero_image_base64 and hero_image_filename:
                        # Decode and save base64 image
                        try:
                            # Decode base64
                            image_data = base64.b64decode(hero_image_base64)
                            
                            # Save image
                            filename = secure_filename(hero_image_filename)
                            filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
                            
                            with open(filepath, 'wb') as f:
                                f.write(image_data)
                            
                            post.hero_image_path = f"/uploads/{filename}"
                            
                        except Exception as e:
                            current_app.logger.error(f"Error importing image for row {row_num}: {e}")
                    
                    # Add to session if new
                    if not existing_post:
                        db.session.add(post)
                    
                except Exception as e:
                    current_app.logger.error(f"Error importing row {row_num}: {e}")
                    skipped_count += 1
                    continue
            
            # Commit all changes
            db.session.commit()
            
            flash(f'Import complete! Imported: {imported_count}, Updated: {updated_count}, Skipped: {skipped_count}', 'success')
            return redirect(url_for('admin.posts_list'))
            
        except Exception as e:
            db.session.rollback()
            current_app.logger.error(f"Import error: {e}")
            flash(f'Import failed: {str(e)}', 'error')
            return redirect(url_for('admin.posts_list'))
    
    # GET request - show import form
    return render_template('admin/posts_import.html')


@admin_bp.route('/prerelease-users', methods=['GET', 'POST'])
@admin_required
def prerelease_users():
    """Manage PreRelease group membership"""
    if request.method == 'POST':
        action = request.form.get('action')
        user_id = request.form.get('user_id')

        user = User.query.get_or_404(user_id)

        if action == 'add':
            if user.is_admin(current_app.config['ADMIN_EMAILS']):
                flash(f'{user.name or user.email} is already an admin with full access', 'warning')
            else:
                user.role = 'prerelease'
                db.session.commit()
                flash(f'{user.name or user.email} added to PreRelease group', 'success')
        elif action == 'remove':
            user.role = 'user'
            db.session.commit()
            flash(f'{user.name or user.email} removed from PreRelease group', 'success')

        return redirect(url_for('admin.prerelease_users'))

    admin_emails = current_app.config['ADMIN_EMAILS']
    prerelease_members = User.query.filter_by(role='prerelease').order_by(User.name).all()
    # Exclude admins — they already have full access
    regular_users = [u for u in User.query.filter_by(role='user').order_by(User.name).all()
                     if u.email not in admin_emails]

    return render_template('admin/prerelease_users.html',
                          prerelease_members=prerelease_members,
                          regular_users=regular_users)


@admin_bp.route('/month-visibility', methods=['GET', 'POST'])
@admin_required
def month_visibility():
    """Manage month visibility (public vs prerelease)"""
    if request.method == 'POST':
        month_key = request.form.get('month_key')
        visibility = request.form.get('visibility')

        if visibility not in ('public', 'prerelease'):
            flash('Invalid visibility value', 'error')
            return redirect(url_for('admin.month_visibility'))

        mv = MonthVisibility.query.filter_by(month_key=month_key).first()
        if mv:
            mv.visibility = visibility
        else:
            mv = MonthVisibility(month_key=month_key, visibility=visibility)
            db.session.add(mv)

        db.session.commit()
        flash(f'{month_key} set to {visibility}', 'success')
        return redirect(url_for('admin.month_visibility'))

    # Get all months that have any posts
    all_months = db.session.query(Post.month_key).distinct().order_by(Post.month_key.desc()).all()
    all_month_keys = [m[0] for m in all_months]

    # Get visibility records
    visibility_records = {mv.month_key: mv.visibility for mv in MonthVisibility.query.all()}

    months_data = []
    for mk in all_month_keys:
        published_count = Post.query.filter_by(month_key=mk, status='published').count()
        draft_count = Post.query.filter_by(month_key=mk, status='draft').count()
        drafts = Post.query.filter_by(month_key=mk, status='draft').all()

        months_data.append({
            'month_key': mk,
            'visibility': visibility_records.get(mk, 'prerelease'),
            'published_count': published_count,
            'draft_count': draft_count,
            'drafts': drafts,
        })

    return render_template('admin/month_visibility.html', months_data=months_data)


@admin_bp.route('/month-visibility/publish-post/<int:post_id>', methods=['POST'])
@admin_required
def publish_post_from_visibility(post_id):
    """Quick publish a draft post from the month visibility page"""
    from datetime import timezone
    post = Post.query.get_or_404(post_id)
    post.status = 'published'
    post.published_at = datetime.now(timezone.utc)
    db.session.commit()
    flash(f'Post "{post.title}" published', 'success')
    return redirect(url_for('admin.month_visibility'))

